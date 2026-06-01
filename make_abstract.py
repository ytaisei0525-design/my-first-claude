# -*- coding: utf-8 -*-
"""
雑誌会要旨（Word）を生成するスクリプト。
- 見本（島田氏の要旨）と同じ形式に揃える。
- 色は使わず、見出し・重要語は太字、補足は（かっこ）で表現。
- 図の位置には空の罫線ボックス（枠）を置き、その下にキャプションを付す。
- 対象論文：Co-assembly of zwitterionic and cationic polymers ...（Biomaterials 2026）
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

JP_FONT = "ＭＳ 明朝"      # 本文（明朝）
JP_FONT_G = "ＭＳ ゴシック"  # 見出し（ゴシック）

doc = Document()

# ---- ページ余白（A4・見本に近い体裁） ----
sec = doc.sections[0]
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)

# ---- 既定スタイル ----
style = doc.styles["Normal"]
style.font.name = JP_FONT
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), JP_FONT)


def set_run_font(run, font=JP_FONT, size=10.5, bold=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)  # 黒のみ（色は使わない）
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)


def add_para(text="", size=10.5, bold=False, align=None, font=JP_FONT,
             space_after=4, space_before=0, line=1.15):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line
    if text:
        r = p.add_run(text)
        set_run_font(r, font=font, size=size, bold=bold)
    return p


def add_body(segments, size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=True):
    """segments: [(text, bold), ...] 太字を混在させた本文段落。"""
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.3
    if first_indent:
        pf.first_line_indent = Pt(10.5)
    for text, bold in segments:
        r = p.add_run(text)
        set_run_font(r, font=JP_FONT, size=size, bold=bold)
    return p


def add_heading(text, size=11.5):
    """【概要】などの大見出し（太字・ゴシック）。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, font=JP_FONT_G, size=size, bold=True)
    return p


def add_subheading(text, size=10.5):
    """結果・考察内の小見出し（太字）。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(7)
    pf.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, font=JP_FONT_G, size=size, bold=True)
    return p


def _set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "8")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        borders.append(e)
    tcPr.append(borders)


def add_figure_box(caption, height_cm=4.2):
    """図を貼るための空の罫線ボックスと、その下のキャプションを追加。"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(16.5)
    cell = table.cell(0, 0)
    cell.width = Cm(16.5)
    _set_cell_border(cell)
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(int(height_cm * 567)))  # cm -> twips
    trHeight.set(qn("w:hRule"), "atLeast")
    trPr.append(trHeight)
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = cp.add_run("（図　ここに貼付）")
    set_run_font(rr, font=JP_FONT, size=9, bold=False)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(8)
    rc = cap.add_run(caption)
    set_run_font(rc, font=JP_FONT, size=9.5, bold=True)


# ============================================================
# ヘッダー部
# ============================================================
add_para("雑誌会要旨", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         font=JP_FONT_G, space_after=2)
add_para("2026/06/02", size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
add_para("M2　＿＿＿＿＿", size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=8)

ptitle = doc.add_paragraph()
ptitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
ptitle.paragraph_format.space_after = Pt(2)
rt = ptitle.add_run(
    "Co-assembly of zwitterionic and cationic polymers for antibacterial "
    "and antithrombotic surfaces via weak electrostatic interactions"
)
set_run_font(rt, font=JP_FONT, size=11.5, bold=True)
add_para("Biomaterials 331 (2026) 124114（J. Wang ら，四川大学 国家生物医学材料工程研究中心）",
         size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

# ============================================================
# 【概要】
# ============================================================
add_heading("【概要】")

add_body([
    ("中心静脈カテーテル、外科用縫合糸、気管挿管チューブ、各種カテーテルなどの埋め込み型・"
     "侵襲性医療デバイスは、集中治療・外科手術・術後管理において不可欠な存在である。しかし、"
     "これらに付随する細菌感染は臨床的安全性を損なう最大の課題の一つであり、臨床データによれば"
     "デバイス関連感染の発生率は", False),
    ("20%を超える", True),
    ("。重度の感染はデバイスの機能不全や再手術を招くだけでなく、", False),
    ("菌血症・敗血症・多臓器不全", True),
    ("といった致命的合併症を引き起こし、患者の死亡率と医療費を大幅に増加させる（図1）。"
     "そのため、デバイス表面を抗菌コーティングで改質して細菌の付着を抑制することは有効な戦略であり、"
     "適切な抗菌剤の選択がその成否を分ける鍵となる。", False),
])

add_body([
    ("既存の抗菌剤のうち、分子鎖上に正電荷を持つ官能基や構造単位を備えた", False),
    ("陽イオン性ポリマー", True),
    ("は、広域スペクトルの抗菌活性を示すことから注目されている。これらは静電的相互作用によって"
     "細菌膜を破壊するため細菌耐性を誘導しにくい。なかでも代表的な", False),
    ("ポリクオタニウム（PQ）", True),
    ("は、第四級アンモニウムカチオンと疎水性の炭素鎖を併せ持ち、カチオンが菌膜へ静電的に作用すると"
     "ともに疎水性炭素鎖が膜の完全性をさらに破壊することで、殺菌効果を相乗的に高める。", False),
])

add_body([
    ("PQをデバイス表面へ被覆する手法としては、ポリマーブラシ法、", False),
    ("層状自己組織化（layer-by-layer：LBL）法", True),
    ("、ドーパミン補助の共有結合グラフト法などがある。なかでもLBL法は、正負に帯電したポリ電解質を"
     "交互に吸着させる簡便かつ汎用性の高い手法として最も広く用いられている。しかし、従来のLBL法では"
     "PQの正電荷が負電荷を持つ", False),
    ("陰イオン性ポリマー（PS）", True),
    ("によって強く遮蔽されるため、PQが菌の細胞膜を「攻撃」する能力が阻害され、抗菌性が制限される。"
     "このため追加の抗菌剤を要することが多い。さらに正負電荷の強い静電引力により電荷と水分子との"
     "結合が減少し、コーティングの親水性が低下して抗ファウリング性能も不十分となる。加えてPQ本来の"
     "正電荷ゆえにタンパク質・細胞・死菌を吸着しやすく、病原体との接触が妨げられて機能が急速に"
     "失われる、という問題があった。", False),
])

add_body([
    ("これに対し著者らは、純粋な陰イオン性ポリマーとは異なる", False),
    ("両性イオン性ポリマー", True),
    ("（ポリスルホベタイン、ポリカルボキシベタイン、ポリホスホコリンなど）に着目した。これらは"
     "分子鎖内に等量の陰イオン（スルホン酸基やリン酸基など）と陽イオン（第四級アンモニウム基など）を"
     "併せ持ち、全体として電気的に中性であるため、非特異的吸着を抑制し、優れた抗ファウリング性・"
     "抗血栓性を示す。両性イオン性ポリマーと陽イオン性ポリマーの間の静電的相互作用は、陰イオン−陽イオン"
     "系のそれよりも", False),
    ("弱い", True),
    ("ことが報告されている。著者らは、この", False),
    ("「弱い静電的相互作用」", True),
    ("が陽イオン基により大きな「自由度」を与えて細菌膜を破壊する能力を保持させると同時に、"
     "荷電基と水分子との結合を維持して緻密な", False),
    ("「水和保護層」", True),
    ("の形成を促し、抗ファウリング性と生体適合性を両立させると仮説を立てた。", False),
])

add_body([
    ("そこで本研究では、両性イオン性ポリマー", False),
    ("PA", True),
    ("（poly[2-(methacryloyloxy)ethyl] dimethyl-(3-sulfopropyl) ammonium hydroxide）と"
     "陽イオン性ポリマー", False),
    ("PQ", True),
    ("をLBL法で組み合わせた複合コーティングを構築した。まず溶液中での両ポリマーの相互作用を吸光度・"
     "粘度・熱変化（ITC）・ゼータ電位・粒子径の観点から解析し、次にコーティングの物理化学的特性を"
     "系統的に評価した。続いて抗菌性・抗ファウリング性・生体適合性をin vitroおよびin vivoで包括的に"
     "検証し、最後に吸収性縫合糸と血液接触カテーテルへ応用して、抗菌性・血液適合性（特に抗血栓性）を"
     "確認した。本研究は、両性イオン性ポリマーと陽イオン性ポリマーを協働させる複合コーティング設計に"
     "理論的・実験的根拠を与えるものである。", False),
])

add_figure_box("図1：弱い静電的相互作用を利用した両性イオン性／陽イオン性ポリマー複合コーティングの概念図",
               height_cm=4.5)

# ============================================================
# 【結果・考察】
# ============================================================
add_heading("【結果・考察】")

add_subheading("溶液中でのPQとPAの相互作用")
add_body([
    ("各ポリマーはフリーラジカル重合により合成した。", False),
    ("¹H NMR", True),
    ("ではモノマー中の二重結合（点a）に対応するピークが重合後に消失し、重合の成功が確認された。"
     "また", False),
    ("XPS", True),
    ("により、PQとPAは窒素（N）の強いシグナルを、PSとPAは硫黄（S）の顕著なシグナルを示した。"
     "PAは純水には不溶であるが、電解質存在下では速やかに溶解性が増す（反ポリ電解質効果）。"
     "そこで塩溶液中での挙動を調べると、PS-PQ・PA-PQ混合物はいずれも混合直後に白濁し、", False),
    ("PS-PQの方がPA-PQより高い濁度", True),
    ("を示した。さらに平衡到達後、PA-PQ混合液は透明化したのに対し、PS-PQには有意な変化が見られなかった。"
     "550 nmの吸光度測定では、PQを1.5当量加えたときPSの吸光度は1.79まで上昇したのに対しPAは1.33に"
     "とどまり、PS-PQがより緻密な凝集体を形成することが示された。平衡後はPS-PQの吸光度がほぼ不変で"
     "あったのに対しPA-PQでは低下し、PA-PQが弱い相互作用による壊れやすい凝集体を作ることが裏付けられた。", False),
])
add_body([
    ("等温滴定型熱量測定（ITC）", True),
    ("では、反対電荷を持つポリマー間の複合体形成が自発的（ΔG＜0）に進行することが示された。"
     "一般に反対電荷物質の凝集は2段階で進む。第1段階は25℃で静電的相互作用に駆動される", False),
    ("エンタルピー律速", True),
    ("過程で、結合点の減少とともにエンタルピー変化が低下する。第2段階は", False),
    ("エントロピー駆動", True),
    ("となる。同濃度でのPS初滴下（PS 32 mM→PQ 4 mM）時のエンタルピー変化はPA-PQよりはるかに大きく、"
     "PA-PQ間の静電的相互作用が非常に弱いことが結論づけられた。なお電荷密度が高いほど反応エントロピー"
     "利得|−TΔS|が大きく、PS-PQの電荷密度がPA-PQより格段に高いことも示された。", False),
])
add_body([
    ("粘度測定では、PS-PQの粘度がPS単独より低下した（分子凝集による）一方、PA-PQの粘度は著しく増加し、"
     "両ポリマー間に", False),
    ("「ゲル状」の架橋ネットワーク", True),
    ("が形成されたことが示された。蛍光分子を結合させたPQ-AFを用いた観察では、同UV吸光度（同濃度）で"
     "PS-PQの平均蛍光強度が34.25であったのに対しPA-PQは43.59を示し、凝集誘起消光の度合いから"
     "PS-PQの方が強固な複合体を作ることが確認された。動的光散乱（DLS）では、陽イオン濃度の増加に伴い"
     "ゼータ電位が負側へ変化した後、PS-PQでは10:30の比で正へ転じ（", False),
    ("電荷反転", True),
    ("）、粒子径も増大した。PA-PQではPAの水和により水分子を取り込んで粒子径が大きくなる一方、"
     "弱い相互作用ゆえ凝集体が壊れやすく、平衡前後で粒径が大きく変化した。これらより、PA-PQは"
     "弱い静電的相互作用によって疎な構造を形成すると結論された。", False),
])

add_figure_box("図2：コーティングの自己組織化過程と表面特性")

add_subheading("コーティングの自己組織化過程と表面特性")
add_body([
    ("PQとPAをLBL法で系に導入し、複合コーティングを作製した。", False),
    ("QCM-d", True),
    ("による実時間モニタリング（周波数シフトΔf）では、ポリマーが基板表面へ逐次吸着し、弱く吸着した"
     "分子はPBS洗浄で除去された。PS-PQではPSの吸着速度がPQ層より低く、組み立てが主に陰イオン−陽イオン"
     "相互作用で進むこと、また吸着が層内で止まらず層間で重なり合うことが示された。PA-PQでは逆にPAの"
     "吸着速度が増加しPQが減少しており、正電荷表面がPA吸着を促進しPQ沈着を抑えることが分かった。", False),
])
add_body([
    ("積層数に対するゼータ電位（図2C）では、PQを最上層にすると", False),
    ("PA-PQ・PS-PQともに正電位", True),
    ("、PAを最上層にするとPA-PQは正電位を保ち、PSを最上層にするとPS-PQは負電位となり、想定どおり"
     "PA（またはPS）と陽イオン性PQがPLA表面へ良好に組み立てられたことが確認された。各層は明確に"
     "分離せず相互に分散・浸透していた。XPSおよびSEM-EDSでは両コーティングで窒素（N）・硫黄（S）の"
     "特性シグナルが検出され、形成が裏付けられた。", False),
])
add_body([
    ("AFM", True),
    ("による粗さ評価では、純PLAは平坦（Rq＝1.06 nm、Ra＝0.67 nm）であったのに対し、複合コーティング"
     "では粗さが増加した。PS-PQはPQの正電荷の不均一分布に由来する典型的な", False),
    ("「島状構造」", True),
    ("を形成し、積層に伴い島の高さ・半径が増大した。一方PA-PQは、PA中の正負電荷がPQとの相互作用に加えて"
     "後続PQ層へのアンカー点を増やすため、積層数の増加に伴い粗さが連続的に減少した。接触角測定では、"
     "PLA（88.73°）に対しPA2.0・PA2.5はそれぞれ", False),
    ("53.3°・21.58°", True),
    ("と優れた親水性を示し、PAの強い水和能が反映された（PS-PQでは正負電荷の相殺で水和能が損なわれ"
     "ほとんど変化なし）。QCM-dを用いたエタノール置換試験では、PA-PQが吸着した液体に伴うΔF・ΔD"
     "（559 Hz・156 ppm）がPS-PQ（225 Hz・89 ppm）より大きく、PA-PQがより多くの液体を吸着し、"
     "より柔らかく", False),
    ("疎で水和能の高い構造", True),
    ("（＝より弱い静電的相互作用）を持つことが示された。PQ-AFを用いた蛍光強度の減衰追跡では、"
     "PA2.5は初日の減衰こそ速いものの長期安定性が高く、PBS浸漬10日でも完全には失活しなかった。", False),
])

add_figure_box("図3：コーティングの抗ファウリング性能と抗血栓性")

add_subheading("抗ファウリング（抗付着）性能")
add_body([
    ("病原体の宿主構造（タンパク質や細胞）への付着は感染症発症の重要な初期段階である。FITC-BSA"
     "（タンパク質）、L929細胞、血小板の付着試験および全血循環試験により抗付着能を評価したところ、"
     "PLAやPS-PQと比較して", False),
    ("PA-PQでは各種生体分子の吸着が有意に減少", True),
    ("し、高い親水性に起因する優れた抗ファウリング性能が示された。SEM観察でも、PAが血小板の活性化と"
     "赤血球・細胞の付着を効果的に抑制することが確認された。", False),
])
add_body([
    ("抗血栓性は、ウサギの", False),
    ("動静脈シャント（ex vivo）モデル", True),
    ("で評価した。1時間循環後、PLA・PS2.0・PS2.5では明らかな血栓沈着（表面被覆率", False),
    ("99.0%・97.1%・89.7%", True),
    ("、質量増加152.9%・90.3%・82.2%）が観察された。血栓は主に凝集タンパク質・赤血球・血小板から"
     "なっていた。これに対しPA-PQ複合コーティングでは血栓被覆と質量増加が大幅に減少し、特に", False),
    ("PA2.5ではほとんど血栓が形成されず（被覆率わずか3.1%）", True),
    ("、PA2.0でも30.3%にとどまった。質量増加（52.5%・51.7%）も主に洗浄時の残留生理食塩水によるもので"
     "あった。優れた親水性がタンパク質吸着や血液・細胞付着を効果的に抑制したことが示された。", False),
])

add_figure_box("図4：コーティングの抗菌性能")

add_subheading("抗菌性能")
add_body([
    ("グラム陽性菌の代表である", False),
    ("S. aureus", True),
    ("をモデル菌とした。微量液体希釈法によるMIC測定では、PS/PQ-M2/0・PA/PQ-M2/0は阻害効果を示さな"
     "かったが、PA/PQはPQ濃度の増加に応じてMIC値が", False),
    ("1/4 M→1/8 M→1/16 M→1/32 M", True),
    ("と段階的に低下し、弱い静電的相互作用が抗菌効果に正に寄与することが示された。PQ濃度の増加で"
     "ゼータ電位の絶対値が小さくなる（＝静電的相互作用が増す）につれMICが倍数的に低下した一方、"
     "2/2・2/3 mg mL⁻¹では弱い結合点が限られMICは頭打ちとなった。対照的にPS/PQでは", False),
    ("PS/PQ-M2/3のみ", True),
    ("が抗菌効果（MIC＝1/32 M）を示し、単純な正負電荷の相互作用は抗菌効果を発揮せず、むしろ正電荷を"
     "遮蔽すること、正電荷が過剰なときのみ抗菌効果が突発的に現れることが示された。", False),
])
add_body([
    ("接触殺菌試験では、PLA表面の生菌数が", False),
    ("6.87×10³ CFU cm⁻²", True),
    ("であったのに対し、PA2.0・PA5.0表面では生菌が検出されず、PA2.5・PA5.5も97.66%・97.62%と高い"
     "抗菌率を示した。PS-PQではPS2.0・PS5.0のみ抗菌効果（97.41%・97.04%）を示し、PS2.5・PS5.5は"
     "PSによるPQの遮蔽で失活した。CLSM観察ではPLA・PS2.0・PS2.5に多数の菌が付着した一方、PA2.0・PA2.5"
     "では著しく少なく、PAの高親水性・抗ファウリング性が反映された。SEMでは菌が不規則な形態と", False),
    ("細胞壁の破裂", True),
    ("を示し、PQが細胞壁の完全性を破壊して殺菌することが確認された。弱い静電的相互作用によりPA-PQ中の"
     "PQとPAは高い自由度を持ち、", False),
    ("PQによる殺菌とPAによる付着抑制という二重機構", True),
    ("を発揮する。これは血液接触デバイスの表面改質において重要な意義を持つ。", False),
])

add_figure_box("図5：コーティングの生体適合性（溶血・細胞毒性）")

add_subheading("生体適合性")
add_body([
    ("埋め込み型デバイスへの応用には良好な生体適合性が不可欠であり、血液適合性と細胞適合性の両面が"
     "重要である。全サンプルの", False),
    ("溶血率は5%未満", True),
    ("であり、赤血球への有意な損傷がないことが示された。特に2.5コーティングの溶血率は2.0コーティング"
     "より低く、PSおよびPAがPQの赤血球への潜在的損傷を遮蔽することが示唆され、接触殺菌の傾向とも"
     "一致した。複合コーティング抽出液を用いた細胞毒性試験では、L929線維芽細胞の生存率が全群で", False),
    ("80%超", True),
    ("となり、細胞毒性がないことが確認された。", False),
])

add_figure_box("図6：ラット皮下埋め込みモデルにおけるin vivo抗菌性能と組織適合性の評価")

add_subheading("生体内（in vivo）での生体適合性と抗菌性能")
add_body([
    ("S. aureusを表面に付着させたサンプルをラットに皮下埋め込みし、抗菌性能を評価した。埋め込み"
     "1日後・3日後の生菌数測定では、PLA・PS2.5表面で有意な細菌生存が認められた一方、", False),
    ("PA2.5では細菌除去率が99%超", True),
    ("で、3日後にはさらに抗菌効果が増強した。PLA・PS2.5表面にはPA2.5より多くの付着物が見られ、"
     "PA2.5が抗菌性に加え優れた抗ファウリング性を持つことが示された。H&E染色では、PLA・PS2.5周囲に"
     "顕著な", False),
    ("炎症細胞浸潤（急性炎症反応）", True),
    ("が観察されたのに対し、PA2.5群は軽度の組織反応にとどまった。CD3（T細胞）・CD68（マクロファージ）の"
     "免疫蛍光染色でもPA2.5の浸潤は有意に少なく、優れた抗菌性により感染関連の炎症が軽減され、"
     "より穏やかな免疫応答と良好な生体適合性を示すことが裏付けられた。", False),
])
add_body([
    ("さらに", False),
    ("ウサギ頸静脈埋め込みモデル", True),
    ("では、S. aureus付着PUカテーテルを頸静脈に埋め込んだ。1日後・3日後の生菌数測定でPA2.5の細菌"
     "除去率は99%超で、肉眼観察・SEMでもPLA・PS2.5に顕著な血栓形成が見られたのに対しPA2.5表面の"
     "血栓被覆は最小限であった。H&E染色では急性炎症反応がPA2.5で大きく軽減し、CD31（新生血管）・"
     "CD68の発現もPLA・PS2.5で強くPA2.5では弱かった。これらより、弱い静電的相互作用によってPQの"
     "毒性が低減されつつ抗菌性が保持され、PAが抗ファウリング・抗炎症効果を発揮することが示された。", False),
])

add_figure_box("図7：ウサギ頸静脈埋め込みモデルにおける抗菌性能と組織反応の評価")

add_subheading("吸収性縫合糸への応用（ラット大腿筋損傷モデル）")
add_body([
    ("複合コーティングを", False),
    ("吸収性外科縫合糸（PGLA）", True),
    ("へ適用し、ラットの大腿筋創傷の治療に用いた。S. aureusを付着させた縫合糸で創を縫合したところ、"
     "1日後・3日後ともにPA2.5表面の生菌はごくわずかで、", False),
    ("殺菌率はいずれも99%超", True),
    ("、3日後にはさらに抗菌効果が増強した。これはコーティング中のPQの抗菌能が、弱い静電的相互作用の"
     "形成によって保持されていることを示す。H&E染色では、PA2.5周囲の炎症細胞浸潤がPGLA・PS2.5より"
     "顕著に少なく、CD3・CD68の免疫蛍光染色でも同様に浸潤が減少しており、優れた抗菌性により感染関連"
     "炎症が軽減され、良好な生体適合性が示された。", False),
])

add_figure_box("図8：ラット大腿筋損傷モデルにおける抗菌性能と組織反応の評価")

# ============================================================
# 【結論】
# ============================================================
add_heading("【結論】")
add_body([
    ("本研究では、陽イオン性ポリマーPQと両性イオン性ポリマーPAをLBL法で層状に自己組織化させることにより、", False),
    ("抗菌性と抗ファウリング性を兼ね備えた複合コーティングの構築に成功", True),
    ("した。複合コーティング中のPQは広域スペクトルの抗菌活性を示し、PAはその強い水和能により緻密な"
     "水和層を形成して、細菌・細胞・タンパク質の非特異的付着を効果的に抑制した。PAとPQの間の", False),
    ("弱い静電的相互作用", True),
    ("はITC・DLS・MICなどの実験により裏付けられ、この相互作用がPQの毒性を大幅に低減しつつ抗菌効果を"
     "概ね維持することを明らかにした。in vitro実験では優れた抗ファウリング性・抗菌性・安定性を、"
     "in vivo実験（ラット皮下・ウサギ頸静脈・ラット大腿筋）では優れた抗菌性・抗ファウリング性・"
     "生体適合性を示し、", False),
    ("埋め込み型デバイス関連感染の予防に大きな可能性", True),
    ("を持つことが示された。今後は、イオン強度がポリマーに与える影響をさらに検討するとともに、"
     "ポリマー比を調整して広域スペクトルの抗菌性能を確保することが課題である。本研究は、両性イオン性"
     "ポリマーと陽イオン性ポリマーの相互作用に対する理解を深める新たな表面改質戦略を提供する。", False),
])

# ============================================================
# 【実験方法】
# ============================================================
add_heading("【実験方法】")

add_subheading("試薬・動物")
add_body([
    ("DMAEMA、AIBME、1-ブロモオクタン、ドーパミン塩酸塩（DOPA）、分岐ポリエチレンイミン（PEI、Mw 25,000）は"
     "Sigma-Aldrich製、SMPSおよびSBMAはAladdin製を用いた。S. aureus（ATCC 25923）、培地・寒天、細菌"
     "生存率アッセイキット、CCK-8等を使用した。動物はニュージーランド白色ウサギ（雄、2〜3 kg）と"
     "Sprague-Dawleyラット（雄、200〜250 g）を用いた。", False),
], first_indent=False)

add_subheading("ポリマーの合成")
add_body([
    ("第四級アンモニウム塩（QAS）モノマーは文献の方法に従って合成した。両性イオン性モノマー（SBMA）、"
     "陽イオン性モノマー（QAS-8C）、陰イオン性モノマー（SMPS）を適切な溶媒（SBMAとSMPSはトリフルオロ"
     "エタノール、QASはイソプロパノール）に溶解し、開始剤AIBME（モノマーの1質量%）を加えて65℃・"
     "無酸素雰囲気下で12時間フリーラジカル重合した。得られた濃縮液を濃縮・洗浄・乾燥し、粉末固体として"
     "ポリマーPA（poly-SBMA）、PQ（poly-QAS-8C）、PS（poly-SMPS）を得た。", False),
], first_indent=False)

add_subheading("溶液中での相互作用解析")
add_body([
    ("25℃でPEAQ-ITC等温滴定型熱量計を用いて、PS-PQおよびPA-PQの相互作用に伴う熱変化を測定した。"
     "各ポリマー溶液（いずれも2 mg mL⁻¹）の粘度はMCR302レオメーター（コーンプレート、直径25 mm、"
     "ギャップ1.0 mm、25℃）で測定した。UVスペクトルは紫外可視分光光度計で、凝集体は倒立蛍光顕微鏡で"
     "観察し、混合後のゼータ電位および粒子径は動的光散乱（DLS）で測定した。", False),
], first_indent=False)

add_subheading("複合コーティングの作製")
add_body([
    ("コーティングはLBL法で作製した（8〜15層、最初の4層を基層とする）。DOPAはトリス緩衝液（pH 8.50）に"
     "2 mg mL⁻¹で、PEI・PA・PSはPBS（pH 7.40）に2 mg mL⁻¹で、PQはPBSに3 mg mL⁻¹で溶解した。"
     "ポリ乳酸（PLA）シートを水・エタノールで超音波洗浄後、DOPA溶液に2時間浸漬して前処理し、PEI・PS・PQ"
     "溶液へ順次10分ずつ浸漬して基層接着層を形成した。続いてPS/PA溶液とPQ溶液へ交互に10分ずつ浸漬し、"
     "各浸漬の間に脱イオン水で洗浄して緩く吸着したポリマーを除去した。40℃で乾燥後、乾燥条件下で保存した。"
     "陽イオン性を最上層とする8層・14層をPA2.0・PA5.0（PS2.0・PS5.0）、両性／陰イオン性を最上層とする"
     "9層・15層をPA2.5・PA5.5（PS2.5・PS5.5）と命名した。同手法をポリウレタン（PU）チューブおよび"
     "吸収性縫合糸（PGLA）にも適用した。", False),
], first_indent=False)

add_subheading("コーティング特性評価")
add_body([
    ("表面ゼータ電位はSurPASS 3（pH 7.40、1 mM KCl）で測定し、組み立て過程はQCM-d（QSense Analyzer）で"
     "周波数変化ΔFを記録して解析した。表面粗さはAFM（Bruker Dimension Icon）で測定した。コーティングの"
     "水和挙動はQCM-dにより周波数（ΔF）と散逸（ΔD）を同時に追跡して評価した（無水・50%・90%エタノールの"
     "交互導入）。", False),
], first_indent=False)

add_subheading("全血循環試験")
add_body([
    ("PLA・PS2.0・PS2.5・PA2.0・PA2.5サンプルをカテーテルに入れ、両端をウサギ全血を充填した遠心管に"
     "接続して体外血液循環ループを構築した。蠕動ポンプで2時間循環させた後、サンプルを洗浄・固定し、"
     "SEMで観察するとともにFDA（緑）染色と蛍光顕微鏡で表面付着を評価した。", False),
], first_indent=False)

add_subheading("抗菌性評価")
add_body([
    ("S. aureusに対するPS/PQおよびPA/PQ混合物のMICを微量液体希釈法で測定した（初期質量濃度比M＝"
     "2/0、2/0.2、2/0.5、2/1、2/2、2/3 mg mL⁻¹）。接触殺菌試験では、菌懸濁液（1×10⁶ CFU mL⁻¹）20 µLを"
     "サンプル表面（1 cm×1 cm）に滴下し、蒸発防止のためポリエチレン膜で覆って37℃で6時間培養した。"
     "その後3 mLの滅菌PBS中で1分間超音波処理して生菌を剥離し、200 µLをMueller-Hinton寒天培地に播種して"
     "37℃で12時間培養後にコロニー数を計数した。SEMによる菌の形態観察、LIVE/DEAD染色とCLSMによる"
     "生死判定も行った。", False),
], first_indent=False)

add_subheading("生体適合性評価")
add_body([
    ("血液適合性は、コーティングを5%ウサギ赤血球と2時間インキュベートして溶血率を求めて評価した"
     "（陽性・陰性対照を設定）。細胞適合性は、複合コーティング抽出液中でL929線維芽細胞を培養し、"
     "FDA染色とCCK-8により細胞生存率を定量して評価した。", False),
], first_indent=False)

add_subheading("in vivo動物試験")
add_body([
    ("全ての動物実験は中国動物保護委員会および四川大学の倫理基準（倫理番号KS2022756）に従い、"
     "1週間の馴化後に実施した。ラット筋肉縫合モデルでは、ペントバルビタール麻酔下で大腿内側筋に"
     "10 mmの切開を加え、あらかじめS. aureus懸濁液（10⁵ CFU mL⁻¹、200 µL、37℃で12時間）に浸漬した"
     "PGLA・PS2.5・PA2.5縫合糸（長さ1 cm）を用いた。手術1日後・3日後にラットを安楽死させ、周囲組織"
     "とともにサンプルを回収して平板計数法で生菌を計数し、周囲軟組織はパラホルムアルデヒドで固定して"
     "病理切片で炎症反応を評価した。ウサギ頸静脈埋め込みモデルおよびラット皮下埋め込みモデルも同様に"
     "評価し、H&E染色・CD3/CD68/CD31免疫蛍光染色を併用した。", False),
], first_indent=False)

add_subheading("統計解析")
add_body([
    ("血液適合性・細胞・in vitro・in vivoの全試験は3回以上反復し、結果は平均±標準偏差（SD）で示した。"
     "GraphPad Prismを用いて一元配置分散分析とLSD（最小有意差）検定を行い、*p＜0.05、**p＜0.01、"
     "***p＜0.001を統計的に有意とした。", False),
], first_indent=False)

out = "/home/user/my-first-claude/雑誌会要旨_PA-PQコーティング.docx"
doc.save(out)
print("保存しました:", out)
