#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""承認済みレイアウトをA0縦(841×1189mm)の実寸PowerPointポスターとして生成する。
   図は後から差し込めるようプレースホルダ枠を配置する。"""

from pptx import Presentation
from pptx.util import Mm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- 配色 ----
C_HEADER = RGBColor(0x0d, 0x3b, 0x66)
C_INTRO  = RGBColor(0xe8, 0xee, 0xf5)
C_INTRO_BAR = RGBColor(0x4a, 0x61, 0x78)
C_MAT    = RGBColor(0xfd, 0xf0, 0xd5)
C_MAT_BAR = RGBColor(0xc9, 0x96, 0x2a)
C_VITRO  = RGBColor(0xdb, 0xea, 0xfe)
C_VITRO_BAR = RGBColor(0x25, 0x63, 0xeb)
C_VIVO   = RGBColor(0xff, 0xe3, 0xd3)
C_VIVO_BAR = RGBColor(0xea, 0x58, 0x0c)
C_CONC   = RGBColor(0xd8, 0xf3, 0xdc)
C_CONC_BAR = RGBColor(0x2d, 0x6a, 0x4f)
C_SEC    = RGBColor(0x33, 0x50, 0x6b)
C_BORDER = RGBColor(0x90, 0xa4, 0xb5)
C_FIGB   = RGBColor(0xb0, 0xbc, 0xc7)
C_WHITE  = RGBColor(0xff, 0xff, 0xff)
C_DARK   = RGBColor(0x22, 0x22, 0x22)
C_BG     = RGBColor(0xf4, 0xf6, 0xf8)
FONT = "Meiryo"  # 日本語フォント(PowerPoint側でレンダリング)

prs = Presentation()
prs.slide_width = Mm(841)
prs.slide_height = Mm(1189)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # 白紙

# 背景
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Mm(841), Mm(1189))
bg.fill.solid(); bg.fill.fore_color.rgb = C_BG; bg.line.fill.background()
bg.shadow.inherit = False

def no_shadow(sp):
    sp.shadow.inherit = False

def box(x, y, w, h, fill, line=C_BORDER, line_w=0.75, rounded=True):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shp_type, Mm(x), Mm(y), Mm(w), Mm(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    no_shadow(sp)
    # 角丸を小さめに
    if rounded:
        try:
            sp.adjustments[0] = 0.04
        except Exception:
            pass
    return sp

def textbox(x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            space_after=2.0, wrap=True):
    """runs: [(text, size, color, bold), ...] 各要素が1段落"""
    tb = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = Mm(2); tf.margin_right = Mm(2)
    tf.margin_top = Mm(1); tf.margin_bottom = Mm(1)
    for i, (txt, size, color, bold) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after); p.space_before = Pt(0)
        r = p.add_run(); r.text = txt
        f = r.font
        f.size = Pt(size); f.bold = bold; f.name = FONT
        f.color.rgb = color
        # 東アジアフォント指定
        rPr = r._r.get_or_add_rPr()
        ea = rPr.makeelement(qn('a:ea'), {'typeface': FONT})
        rPr.append(ea)
    return tb

def figph(x, y, w, h, caption):
    """図プレースホルダ(灰色の枠 + キャプション)。差し込み位置の目印。"""
    sp = box(x, y, w, h, RGBColor(0xf0,0xf3,0xf6), line=C_FIGB, line_w=1.0, rounded=False)
    textbox(x, y, w, h,
            [("【図を差し込み】", 11, C_FIGB, True),
             (caption, 11, RGBColor(0x6b,0x78,0x84), False)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=2)

def panel(x, y, w, h, title, bullets, fill, bar, fig_caption=None, fig_h=0,
          body_size=12, title_size=14):
    box(x, y, w, h, fill, line=C_BORDER, line_w=1.0)
    # タイトルバー
    bar_sp = box(x, y, w, 9, bar, line=None, rounded=False)
    textbox(x+1, y, w-2, 9,
            [(title, title_size, C_WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
    # 本文
    body_h = h - 11 - (fig_h+3 if fig_caption else 0)
    runs = [(b, body_size, C_DARK, False) for b in bullets]
    textbox(x+1.5, y+10, w-3, body_h, runs, space_after=2.0)
    # 図
    if fig_caption:
        figph(x+3, y+h-fig_h-3, w-6, fig_h, fig_caption)

def sectionbar(x, y, w, label, color, size=18):
    box(x, y, w, 12, color, line=None, rounded=False)
    textbox(x+4, y, w-8, 12, [(label, size, C_WHITE, True)],
            anchor=MSO_ANCHOR.MIDDLE)

# ===================== レイアウト =====================
M = 8
gap = 5
cw = 841 - 2*M

# --- ヘッダー ---
hh = 62
box(M, M, cw, hh, C_HEADER, line=None)
textbox(M, M+5, cw, 40,
        [("ポリグリセロールデンドリマーとグリコールキトサンの水素結合による", 26, C_WHITE, True),
         ("注入可能・自己修復型 超分子ハイドロゲルの創製", 26, C_WHITE, True)],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, space_after=2)
textbox(M, M+46, cw, 12,
        [("〇山道 大誠¹  Cho Iksung²³  田中 賢²³ （九州大学 工学部／工学府／先導物質化学研究所）",
          14, RGBColor(0xcf,0xe0,0xf0), False)],
        align=PP_ALIGN.CENTER)

# --- Introduction(全幅・上部)---
iy = M + hh + gap
sectionbar(M, iy, cw, "Introduction ｜ 背景・着眼点・目的", C_HEADER, size=20)
sub_y = iy + 14
sub_h = 138
sub_gap = 5
sub_w = (cw - 2*sub_gap) / 3
sx0 = M; sx1 = M + sub_w + sub_gap; sx2 = M + 2*(sub_w+sub_gap)
panel(sx0, sub_y, sub_w, sub_h, "化学療法の課題",
      ["・全身投与は腫瘍到達率 中央値0.7%",
       "・骨髄抑制・脱毛など重篤な副作用",
       "・PTXは水溶解度 < 0.5 µg/mL",
       "・可溶化剤 Cremophor EL が重篤毒性",
       "→ 無毒可溶化＋局所投与のDDSが必要"],
      C_INTRO, C_INTRO_BAR, fig_caption="図: PTX/CrEL構造・0.7%模式図", fig_h=54)
panel(sx1, sub_y, sub_w, sub_h, "着眼点: 2機能を1材料で",
      ["・PGD: 単分散・中間水・生体適合性",
       "・単分子ハイドロトロープ機能",
       "  → PTXを高可溶化(PEG400の約10倍)",
       "・shear-thinning型: 注入後に自己修復",
       "  → 液だれせず患部に留まる"],
      C_INTRO, C_INTRO_BAR, fig_caption="図: PGD構造 / 溶解度 / shear機構", fig_h=54)
panel(sx2, sub_y, sub_w, sub_h, "目的・戦略",
      ["目的: ハイドロトロープ機能と自己修復性を",
       "両立した注入ゲルの開発",
       "戦略: PGD＋GCの水素結合(物理架橋)で",
       "自発形成する超分子ハイドロゲルを利用",
       "→ PGDがPTXを高濃度保持し局所徐放"],
      C_INTRO, C_INTRO_BAR, fig_caption="図: PGD＋GC 超分子ゲル形成(中心図)", fig_h=54)

# --- Results & Discussion ---
ry = sub_y + sub_h + gap
sectionbar(M, ry, cw, "Results & Discussion ｜ データ（左→右で研究の流れ）", C_SEC, size=20)

# 下部は薄い結論帯のみ(In vivo帯は廃止)
band_conc_h = 34
conc_y = 1189 - M - band_conc_h

colw = (cw - 2*gap) / 3
dx0 = M; dx1 = M + colw + gap; dx2 = M + 2*(colw+gap)

# --- 列ごとの見出しバー ---
head_y = ry + 14
head_h = 10
def colhead(x, label, color):
    box(x, head_y, colw, head_h, color, line=None, rounded=False)
    textbox(x, head_y, colw, head_h, [(label, 15, C_WHITE, True)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
colhead(dx0, "材料合成・インジェクタブル特性", C_SEC)
colhead(dx1, "In vitro 生体適合性評価", C_VITRO_BAR)
colhead(dx2, "実用性評価 (In vitro / In vivo)", C_VIVO_BAR)

data_y = head_y + head_h + 4
data_h = conc_y - gap - data_y

# 左列=3枚, 中列=2枚, 右列=2枚
ph3 = (data_h - 2*gap) / 3
ph2 = (data_h - gap) / 2
f3 = ph3 - 40
f2 = ph2 - 50

# ---- 左列: 合成 → 自己修復性 → シアシニング ----
panel(dx0, data_y, colw, ph3, "① PGD G4 の精密合成〔新〕",
      ["・Divergent法(全8段階)で単分散PGD G4を合成",
       "・撹拌強化＋反応時間2倍で高世代を最適化",
       "・収率90%, MALDI単一ピーク [M+Na]⁺=3488"],
      C_MAT, C_MAT_BAR, fig_caption="図: MALDI(最適化前後) / 収率表", fig_h=f3)
panel(dx0, data_y+ph3+gap, colw, ph3, "② 自己修復性 (レオロジー)〔目玉〕",
      ["・せん断除去で G′ が即座に回復",
       "・元の強固なゲル構造へ再構築",
       "・複数サイクルで再現性よく回復"],
      C_VITRO, C_VITRO_BAR, fig_caption="図: ステップ歪みサイクル (G′/G″)", fig_h=f3)
panel(dx0, data_y+2*(ph3+gap), colw, ph3, "③ シアシニング性 (レオロジー)",
      ["・高せん断でゲルが流動化(ゾル化)",
       "・粘度が急低下 → 注射針を通過可能",
       "・大歪500%で G′<G″"],
      C_VITRO, C_VITRO_BAR, fig_caption="図: せん断速度-粘度 / 大歪での G′低下", fig_h=f3)

# ---- 中列: In vitro 安全性 → 分解性 ----
panel(dx1, data_y, colw, ph2, "④ In vitro 安全性｜Live/Dead〔新〕",
      ["・D1細胞を3次元ゲル内に包埋し培養",
       "・1/5/9日とも死細胞はほぼ観察されず",
       "・溶解後も毒性なし",
       "・→ 化学架橋剤なしで高い細胞生存率"],
      C_VITRO, C_VITRO_BAR,
      fig_caption="図: Live/Dead 蛍光像\n(PG3-1 / PG4-025 / PG4-0125 ×1・5・9日)", fig_h=f2)
panel(dx1, data_y+ph2+gap, colw, ph2, "⑤ In vitro 分解性〔新〕",
      ["・37℃ PBS中の重量変化で分解を評価",
       "・PG3-1:3日 / PG4-0125:4日 / PG4-025:25日",
       "・世代・混合比で滞留期間を制御",
       "・→ 数日〜数週間で任意に制御可能"],
      C_VITRO, C_VITRO_BAR, fig_caption="図: 分解曲線 (重量比 vs 時間)", fig_h=f2)

# ---- 右列: 注入試験(小) → In vivo 組織学的評価 H&E(大) ----
ph_rt = (data_h - gap) * 0.36
ph_rb = (data_h - gap) * 0.64
panel(dx2, data_y, colw, ph_rt, "⑥ 注入試験 (In vitro & In vivo)〔新〕",
      ["・in vitro: 赤色化ゲルをシリンジ吐出→即ゲル化",
       "・in vivo: マウス背部皮下へ注入(0.25 mL)",
       "・抵抗なく注入→液漏れせずドーム状デポ形成"],
      C_VIVO, C_VIVO_BAR,
      fig_caption="図: in vitro 注入写真 / in vivo マウス皮下デポ写真", fig_h=ph_rt-44)
panel(dx2, data_y+ph_rt+gap, colw, ph_rb, "⑦ In vivo 組織学的評価 H&E染色〔考察・新〕",
      ["・Day 0/5/14で皮下デポを採取しH&E評価",
       "・ゲル辺縁から宿主細胞が内部へ浸潤(矢印)",
       "・Day5→14で浸潤進行 → 緩やかに生分解・組織置換",
       "・14日後もデポ残存。重度の炎症・壊死なし(高適合性)",
       "・in vitro分解(≈25日)と整合 → 局所貯留に好適"],
      C_VIVO, C_VIVO_BAR,
      fig_caption="図: 実験スキーム / Optical(デポ) / H&E組織像\n(Native vs PG4-025 ×0・5・14日)", fig_h=ph_rb-58)

# --- Conclusion(薄い帯)---
box(M, conc_y, cw, band_conc_h, C_CONC, line=C_CONC_BAR, line_w=1.2)
textbox(M+3, conc_y+2, cw-6, 18,
        [("Conclusion:  単分散PGD G4を確立 → 水素結合で超分子ゲル形成（自己修復・分解性を制御・高生体適合性）"
          " → in vivoで局所デポを形成し実用性を実証", 13, C_CONC_BAR, True)])
textbox(M+3, conc_y+20, cw-6, 18,
        [("Future: PTX徐放定量(HPLC) / 担がんマウスで抗腫瘍効果 / PGD G5合成 / DSCで中間水評価　　"
          "Ref: Ooya Gels 2022 ; Cho & Ooya 2018 ; Yamazaki Langmuir 2021 ほか", 11,
          RGBColor(0x3a,0x4a,0x55), False)])

prs.save("poster/poster_A0.pptx")
print("PowerPoint生成完了: poster/poster_A0.pptx")
